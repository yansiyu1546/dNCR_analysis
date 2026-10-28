"""Modeling/data utilities and standalone prediction CLI. See README.md."""
from __future__ import annotations


import argparse
import csv
import hashlib
import json
from datetime import datetime, timedelta, timezone
from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RUN_DIRECTORIES = ('cache', 'models', 'results', 'imputations', 'figures', 'reports', 'qa_sheets')
RUN_FILES = ('使用说明.txt', '完整分析入口.ipynb', 'SHA256SUMS.txt', 'offline_predictions.csv')


def project_root(script):
    """Private analysis workspace; never search outside this package for code."""
    import os
    package = Path(script).resolve().parent
    configured = os.environ.get('DNCR_ANALYSIS_ROOT')
    root = Path(configured).expanduser().resolve() if configured else package / 'private_analysis'
    # Analysis initialization can archive and replace output folders. Keep it
    # separate from the distributed inference models and source code.
    if root == package or package.is_relative_to(root):
        raise ValueError('Choose a dedicated analysis workspace, not the package root or its parent.')
    for folder in ['model']:
        protected = package / folder
        if root == protected or root.is_relative_to(protected) or protected.is_relative_to(root):
            raise ValueError('The analysis workspace must not overlap the distributed package files.')
    return root


def read_frame(path):
    """Check headers before pandas can rename duplicate columns to .1, .2, etc."""
    path = Path(path)
    if path.suffix.lower() == '.xlsx':
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            headers = list(next(book.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True)))
        finally:
            book.close()
        while headers and headers[-1] is None:
            headers.pop()
        frame = pd.read_excel(path, sheet_name=0)
    elif path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as stream:
            headers = next(csv.reader(stream))
        frame = pd.read_csv(path)
    else:
        raise ValueError(f'Unsupported input file: {path.name}; use .xlsx or .csv.')
    normalized = [str(h).strip() if h is not None else '' for h in headers]
    duplicates = sorted({h for h in normalized if normalized.count(h) > 1})
    if duplicates:
        raise ValueError(f'{path.name} has duplicate column names: {duplicates}. Keep one source column per variable.')
    if not normalized or '' in normalized or len(frame.columns) != len(normalized):
        raise ValueError(f'{path.name} contains an empty or inconsistent column header.')
    frame.columns = normalized
    if frame.empty:
        raise ValueError(f'{path.name} contains no data rows.')
    return frame


def numeric_columns(frame, columns, source):
    frame = frame.copy()
    for column in columns:
        original = frame[column].replace(r'^\s*$', np.nan, regex=True)
        converted = pd.to_numeric(original, errors='coerce')
        bad = original.notna() & converted.isna()
        if bad.any():
            excel_rows = (np.flatnonzero(bad.to_numpy())[:8] + 2).tolist()
            raise ValueError(f'{source}: nonnumeric {column} values at Excel rows {excel_rows}.')
        if np.isinf(converted.to_numpy(dtype=float)).any():
            raise ValueError(f'{source}: {column} contains infinite values.')
        frame[column] = converted
    return frame


def candidate_columns(frame, config):
    required = set(config.get('required_predictors', []))
    excluded = set(config.get('exclude', []))
    missing = sorted(required - set(frame.columns))
    if missing:
        raise ValueError(f'Missing required candidate predictors: {missing}.')
    if required & excluded:
        raise ValueError(f'Required predictors must participate in selection: {sorted(required & excluded)}.')
    return [c for c in frame.columns if c != 'dNCR' and c not in excluded]


def load_cohorts(root, config):
    root = Path(root)
    cohorts = []
    for filename in ('training.xlsx', 'external.xlsx'):
        frame = read_frame(root / 'data' / filename)
        if 'dNCR' not in frame:
            raise ValueError(f'{filename} is missing the dNCR outcome column.')
        candidate_columns(frame, config)
        frame = numeric_columns(frame, list(frame.columns), filename)
        if frame.dNCR.isna().any() or set(frame.dNCR.unique()) != {0, 1}:
            raise ValueError(f'{filename}: dNCR must contain both 0 and 1 with no missing values.')
        frame['dNCR'] = frame.dNCR.astype(int)
        cohorts.append(frame)
    training, external = cohorts
    missing = sorted(set(training.columns) - set(external.columns))
    extra = sorted(set(external.columns) - set(training.columns))
    if missing or extra:
        raise ValueError(f'External cohort columns differ: missing={missing}, extra={extra}.')
    external = external.loc[:, training.columns].copy()
    empty = [c for c in candidate_columns(training, config) if training[c].isna().all()]
    if empty:
        raise ValueError(f'All development values are missing for predictors: {empty}.')
    return training, external


def analysis_fingerprint(training, external, config):
    content = training.to_csv(index=False) + external.to_csv(index=False) + json.dumps(config, sort_keys=True)
    return hashlib.sha256(content.encode()).hexdigest()


def source_hashes(root):
    return {name: hashlib.sha256((Path(root) / 'data' / name).read_bytes()).hexdigest()
            for name in ('training.xlsx', 'external.xlsx')}


def _archive_directory_contents(source, destination, archive, moved):
    """A Windows directory handle can block renaming while its files remain movable."""
    source_root = source.resolve()
    archive_root = archive.resolve()
    directories = [destination]
    files = []
    # Validate the complete plan before moving any files in this directory.
    for entry in sorted(source.rglob('*')):
        if (entry.is_symlink() or entry.is_junction()
                or not entry.resolve().is_relative_to(source_root)):
            raise ValueError(f'Will not archive a linked path: {entry}')
        target = destination / entry.relative_to(source)
        if not target.resolve().is_relative_to(archive_root):
            raise ValueError(f'Archive target is outside the archive directory: {target}')
        if entry.is_dir():
            directories.append(target)
        elif entry.is_file():
            files.append((entry, target))
        else:
            raise ValueError(f'Unsupported archive entry: {entry}')
    for directory in directories:
        directory.mkdir(parents=True, exist_ok=True)
    print(f'Archiving files while keeping the active directory: {source.name}', flush=True)
    for entry, target in files:
        entry.rename(target)
        moved.append((entry, target))


def initialize_run(root, fingerprint, fresh=False):
    """Explicit --fresh archives prior artifacts; ordinary runs retain resume support."""
    root = Path(root).resolve()
    marker = root / 'cache' / 'fingerprint.txt'
    audit_path = root / 'results' / 'input_audit.json'
    previous = marker.read_text(encoding='utf-8').strip() if marker.exists() else None
    if previous is None and audit_path.exists():
        previous = json.loads(audit_path.read_text(encoding='utf-8')).get('input_sha256')
    if fresh:
        archive_path = root / 'archive'
        archive_base = archive_path.resolve()
        if archive_base.parent != root or archive_path.is_symlink() or archive_path.is_junction():
            raise ValueError('The archive directory must stay inside the project root.')
        sources = []
        for name in RUN_DIRECTORIES + RUN_FILES:
            source = root / name
            if source.exists():
                if source.is_symlink() or source.is_junction() or source.resolve().parent != root:
                    raise ValueError(f'Will not archive a path outside the project: {source}')
                sources.append(source)
        if sources:
            stamp = datetime.now(timezone(timedelta(hours=8))).strftime('%Y%m%d_%H%M%S')
            archive = archive_base / f'{stamp}_{uuid4().hex[:8]}'
            archive.mkdir(parents=True, exist_ok=False)
            moved = []
            try:
                for source in sources:
                    destination = archive / source.name
                    try:
                        source.rename(destination)
                    except PermissionError:
                        if not source.is_dir():
                            raise
                        _archive_directory_contents(source, destination, archive, moved)
                    else:
                        moved.append((source, destination))
            except BaseException as error:
                rollback_errors = []
                for source, destination in reversed(moved):
                    try:
                        if source.exists():
                            raise FileExistsError(f'Cannot restore over an existing path: {source}')
                        destination.rename(source)
                    except OSError as rollback_error:
                        rollback_errors.append(str(rollback_error))
                if rollback_errors:
                    raise RuntimeError(f'Archiving failed; some prior files remain in {archive}. '
                                       f'Restore these before starting a new run: {rollback_errors}') from error
                raise
            print(f'Previous run archived to: {archive}', flush=True)
    elif previous is not None and previous != fingerprint:
        raise RuntimeError('Input data, candidate columns or analysis configuration changed. '
                           'Run step_2_run_primary_analysis.py --jobs 6 --fresh to archive previous artifacts and rebuild.')
    elif previous is None and any((root / 'cache').glob('*.joblib')):
        raise RuntimeError('Unversioned preprocessing caches found. Use --fresh before a new analysis.')
    for name in RUN_DIRECTORIES:
        (root / name).mkdir(parents=True, exist_ok=True)
    (root / 'results' / 'analysis_complete.json').unlink(missing_ok=True)
    marker.write_text(fingerprint, encoding='utf-8')


def require_current_analysis(root, config, require_complete=True):
    """Read-only gate for evaluation, explanations, prediction, and report generation."""
    root = Path(root)
    training, external = load_cohorts(root, config)
    fingerprint = analysis_fingerprint(training, external, config)
    audit_path = root / 'results' / 'input_audit.json'
    if not audit_path.is_file():
        raise RuntimeError('No analysis for the current data. Run step_2_run_primary_analysis.py first.')
    audit = json.loads(audit_path.read_text(encoding='utf-8'))
    columns = candidate_columns(training, config)
    if audit.get('input_sha256') != fingerprint or audit.get('candidate_predictors') != columns:
        raise RuntimeError('Saved results do not match the current data, candidate columns or configuration. '
                           'Run step_2_run_primary_analysis.py --jobs 6 --fresh before this step.')
    marker = root / 'cache' / 'fingerprint.txt'
    if marker.exists() and marker.read_text(encoding='utf-8').strip() != fingerprint:
        raise RuntimeError('Cache fingerprint does not match the current input audit.')
    if require_complete:
        completed = root / 'results' / 'analysis_complete.json'
        if not completed.exists() or json.loads(completed.read_text(encoding='utf-8')).get('input_sha256') != fingerprint:
            raise RuntimeError('The main analysis has not completed for these inputs. Resume step_2_run_primary_analysis.py.')
    return training, external, audit


def require_current_artifact(artifact, audit):
    if artifact.get('input_sha256') != audit['input_sha256']:
        raise RuntimeError('Saved model belongs to a different input dataset; rebuild the main analysis.')
    if artifact['preprocess']['columns'] != audit['candidate_predictors']:
        raise RuntimeError('Saved model candidate columns do not match the current data.')




import copy, json, warnings
from pathlib import Path
import numpy as np
import pandas as pd
from scipy.special import expit, logit
from sklearn.experimental import enable_iterative_imputer  # noqa: F401
from sklearn.impute import IterativeImputer
from sklearn.preprocessing import StandardScaler
from sklearn.ensemble import RandomForestClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.neural_network import MLPClassifier
from sklearn.svm import SVC
from sklearn.dummy import DummyClassifier
from sklearn.metrics import roc_auc_score, roc_curve, brier_score_loss
from sklearn.utils.class_weight import compute_sample_weight
from sklearn.exceptions import ConvergenceWarning
from boruta import BorutaPy
from xgboost import XGBClassifier

CONFIG = dict(seed=20261017, outer_folds=5, inner_folds=3, imputations=5,
              boruta_seeds=2, boruta_iterations=50, boruta_trees=128,
              boruta_depth=5, boruta_percentile=100, stability_cutoff=0.8,
              trials_per_model=12, target_sensitivity=0.85,
              winsor_quantiles=[0.01,0.99], exclude=[], required_predictors=[],
              categorical=['Surgery_Type'],
              pair_priority=[['Operation_Time','Anesthesia_Time'],['MOCA_Score','Education']],
              bootstrap=2000, extra_boruta_seeds=10)
NAMES = ['LR','RF','XGB','ANN','SVM']
SPACES = {
 'LR':dict(C=[0.001,0.01,0.1,1,10,100],l1_ratio=[0.,0.5,1.],class_weight=[None,'balanced']),
 'RF':dict(n_estimators=[200,400],max_depth=[3,5,8,None],min_samples_leaf=[3,5,10],max_features=['sqrt',0.8],class_weight=[None,'balanced']),
 'XGB':dict(n_estimators=[100,200,300],max_depth=[1,2,3],learning_rate=[0.01,0.03,0.05],subsample=[0.7,0.9],colsample_bytree=[0.7,0.9],min_child_weight=[3,5,10],reg_alpha=[0.,0.5,1.],reg_lambda=[2.,5.,10.],scale_pos_weight=[1.,1.5,2.]),
 'ANN':dict(hidden_layer_sizes=[(8,),(16,),(16,8)],activation=['relu','tanh'],alpha=[0.001,0.01,0.1,1.],learning_rate_init=[0.001,0.003],batch_size=[16,32],class_weight=[None,'balanced']),
 'SVM':dict(C=[0.01,0.1,1,10,100],gamma=['scale',0.001,0.01,0.1],class_weight=[None,'balanced'])}

def seed(*parts):
    return int(np.random.SeedSequence([CONFIG['seed'],*map(int,parts)]).generate_state(1)[0])

def jsonwrite(path,obj):
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    Path(path).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=lambda x:x.tolist() if hasattr(x,'tolist') else str(x)),encoding='utf-8')

def encoded(a,cols):
    out=[]; groups=[]; labels=[]
    for j,c in enumerate(cols):
        if c in CONFIG['categorical']:
            for v in [1,2,3]:
                out.append((a[:,j]==v).astype(float));groups.append(c);labels.append(f'Surgery_type_{v}')
        else:
            out.append(a[:,j]);groups.append(c);labels.append(c)
    return np.asarray(out).T,groups,labels

def clip_values(frame,lo,hi):
    return np.clip(frame.to_numpy(dtype=float),lo,hi)

def impute_new(imp,a,random_seed):
    # Copy so predictions never mutate a saved imputer's posterior RNG.
    obj=copy.deepcopy(imp)
    obj.random_state_=np.random.RandomState(random_seed)
    return obj.transform(a)

def boruta_once(a,y,cols,random_seed):
    z,groups,labels=encoded(a,cols)
    forest=RandomForestClassifier(n_estimators=CONFIG['boruta_trees'],max_depth=CONFIG['boruta_depth'],min_samples_leaf=3,class_weight='balanced',n_jobs=1,random_state=random_seed)
    selector=BorutaPy(forest,n_estimators=CONFIG['boruta_trees'],perc=CONFIG['boruta_percentile'],alpha=0.05,two_step=True,max_iter=CONFIG['boruta_iterations'],random_state=random_seed,verbose=0)
    selector.fit(z,y)
    return {c:dict(confirmed=bool(np.any(selector.support_[np.array(groups)==c])),tentative=bool(np.any(selector.support_weak_[np.array(groups)==c]))) for c in cols}

def priority(selected):
    selected=list(selected)
    for keep,drop in CONFIG['pair_priority']:
        if keep in selected and drop in selected:selected.remove(drop)
    return selected

def prepare(X,y,train_idx,valid_idx,tag,context_id,outdir,input_sha256=None):
    """Everything in this function is fitted only on train_idx."""
    import joblib
    outdir=Path(outdir);cache=outdir/'cache'/f'{tag}.joblib'
    cols=candidate_columns(X,CONFIG)
    if cache.exists():
        old=joblib.load(cache)
        if (old.get('columns')!=cols or old.get('input_sha256')!=input_sha256
                or not np.array_equal(old['train_idx'],train_idx)
                or not np.array_equal(old['valid_idx'],valid_idx)):
            raise RuntimeError(f'Stale preprocessing cache: {tag}. Restart the main analysis with --fresh.')
        return str(cache)
    tr=X.iloc[train_idx];va=X.iloc[valid_idx];cols=list(X.columns)
    lo=tr.min().to_numpy(float);hi=tr.max().to_numpy(float)
    for j,c in enumerate(cols):
        if tr[c].nunique()>10 and c not in CONFIG['categorical']:
            lo[j],hi[j]=tr[c].quantile(CONFIG['winsor_quantiles']).to_numpy()
    a=clip_values(tr,lo,hi);b=clip_values(va,lo,hi)
    members=[];votes=[]
    for m in range(CONFIG['imputations']):
        rs=seed(context_id,10,m)
        imp=IterativeImputer(max_iter=10,sample_posterior=True,random_state=rs,min_value=lo,max_value=hi,skip_complete=False,initial_strategy='median')
        with warnings.catch_warnings():
            warnings.simplefilter('ignore',ConvergenceWarning)
            imputed=imp.fit_transform(a)
        new=impute_new(imp,b,seed(context_id,11,m)) if len(b) else np.empty((0,len(cols)))
        members.append(dict(imputer=imp,train=imputed,valid=new,imputation_seed=rs,transform_seed=seed(context_id,11,m)))
        for r in range(CONFIG['boruta_seeds']):
            bs=seed(context_id,20,r)
            res=boruta_once(imputed,y[train_idx],cols,bs)
            for c,v in res.items():votes.append(dict(context=tag,imputation=m+1,boruta_run=r+1,imputation_seed=rs,boruta_seed=bs,variable=c,**v))
    vote=pd.DataFrame(votes)
    freq=vote.groupby('variable').confirmed.mean().reindex(cols)
    selected=priority(freq[freq>=CONFIG['stability_cutoff']].index.tolist())
    bundle=dict(tag=tag,context_id=context_id,columns=cols,lower=lo,upper=hi,members=members,selected=selected,frequencies=freq,train_idx=np.asarray(train_idx),valid_idx=np.asarray(valid_idx),input_sha256=input_sha256)
    joblib.dump(bundle,cache,compress=3)
    vote.to_csv(outdir/'results'/f'boruta_{tag}.csv',index=False)
    print(f'PREP {tag}: n={len(train_idx)}, selected={selected}',flush=True)
    return str(cache)

def matrix(a,cols,selected):
    z,groups,labels=encoded(a,cols)
    mask=np.array([g in selected for g in groups])
    return z[:,mask],[l for l,k in zip(labels,mask) if k]

def model_factory(name,p,rs):
    p=p.copy();cw=p.pop('class_weight',None)
    if name=='LR':return LogisticRegression(solver='saga',max_iter=3000,tol=1e-4,random_state=rs,class_weight=cw,**p),None
    if name=='RF':return RandomForestClassifier(n_jobs=1,random_state=rs,class_weight=cw,**p),None
    if name=='XGB':return XGBClassifier(objective='binary:logistic',eval_metric='logloss',tree_method='hist',n_jobs=1,random_state=rs,**p),None
    if name=='ANN':return MLPClassifier(solver='adam',max_iter=600,early_stopping=True,validation_fraction=0.15,n_iter_no_change=30,tol=1e-4,random_state=rs,**p),cw
    if name=='SVM':return SVC(kernel='rbf',probability=False,random_state=rs,class_weight=cw,**p),None
    raise ValueError(name)

def model_score(model,z,name):
    if name=='SVM' and not isinstance(model,DummyClassifier):return model.decision_function(z)
    return model.predict_proba(z)[:,1]

def fit_members(bundle,y,name,params):
    models=[];preds=[];trainpreds=[];conv=[]
    for m,member in enumerate(bundle['members']):
        a,labels=matrix(member['train'],bundle['columns'],bundle['selected'])
        b,_=matrix(member['valid'],bundle['columns'],bundle['selected'])
        scaler=None
        if a.shape[1]==0:
            a=np.zeros((len(a),1));b=np.zeros((len(b),1));model=DummyClassifier(strategy='prior');cw=None
        else:
            if name in ['LR','ANN','SVM']:
                scaler=StandardScaler().fit(a);a=scaler.transform(a);b=scaler.transform(b) if len(b) else b
            model,cw=model_factory(name,params,seed(bundle['context_id'],30,NAMES.index(name),m))
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter('always',ConvergenceWarning)
            weight=compute_sample_weight('balanced',y) if cw=='balanced' else None
            if weight is None:model.fit(a,y)
            else:model.fit(a,y,sample_weight=weight)
        conv.append(any(issubclass(w.category,ConvergenceWarning) for w in caught))
        models.append(dict(model=model,scaler=scaler,feature_names=labels))
        preds.append(model_score(model,b,name) if len(b) else np.array([]))
        trainpreds.append(model_score(model,a,name))
    return models,np.mean(preds,axis=0),np.mean(trainpreds,axis=0),conv

def calibration_input(raw,name):
    a=np.asarray(raw)
    return a.reshape(-1,1) if name=='SVM' else logit(np.clip(a,1e-6,1-1e-6)).reshape(-1,1)

def fit_calibrator(y,raw,name):
    c=LogisticRegression(C=1e6,solver='lbfgs',max_iter=1000)
    c.fit(calibration_input(raw,name),y)
    return c

def calibrated(c,raw,name):return c.predict_proba(calibration_input(raw,name))[:,1]

def choose_threshold(y,p,target=None):
    target=CONFIG['target_sensitivity'] if target is None else target
    fpr,tpr,thresholds=roc_curve(y,p,drop_intermediate=False)
    valid=np.flatnonzero((tpr>=target)&np.isfinite(thresholds))
    return float(thresholds[valid[0]])

def predict_saved(artifact,frame):
    bundle=artifact['preprocess'];name=artifact['name']
    a=clip_values(frame[bundle['columns']],bundle['lower'],bundle['upper'])
    pp=[]
    for m,(member,fitted) in enumerate(zip(bundle['members'],artifact['fitted'])):
        values=impute_new(member['imputer'],a,member['transform_seed'])
        z,_=matrix(values,bundle['columns'],bundle['selected'])
        if z.shape[1]==0:z=np.zeros((len(z),1))
        if fitted['scaler'] is not None:z=fitted['scaler'].transform(z)
        pp.append(model_score(fitted['model'],z,name))
    return calibrated(artifact['calibrator'],np.mean(pp,axis=0),name)

def predict_complete_selected(artifact,frame):
    """Equivalent inference when every selected predictor is observed.

    Auxiliary imputation columns are unnecessary in this specific case.
    """
    b=artifact['preprocess'];selected=b['selected'];cols=b['columns'];name=artifact['name']
    if frame[selected].isna().any().any():raise ValueError('Selected predictors contain missing values; use the full candidate-input pipeline.')
    fill=(np.asarray(b['auxiliary_fill_values'],dtype=float) if 'auxiliary_fill_values' in b else np.median(b['members'][0]['train'],axis=0))
    a=np.tile(fill,(len(frame),1))
    for c in selected:a[:,cols.index(c)]=frame[c].to_numpy(float)
    a=np.clip(a,b['lower'],b['upper']);z,_=matrix(a,cols,selected);pp=[]
    for member in artifact['fitted']:
        zz=member['scaler'].transform(z) if member['scaler'] is not None else z
        pp.append(model_score(member['model'],zz,name))
    return calibrated(artifact['calibrator'],np.mean(pp,axis=0),name)


import argparse
import json
from pathlib import Path

import joblib
import numpy as np
import pandas as pd

FORMAT = 'dncr-inference-v1'


def predict_frame(artifact, frame):
    if artifact.get('release_format') != FORMAT:
        raise ValueError('Use the exported dncr-inference-v1 models.')
    bundle = artifact['preprocess']
    selected, columns = bundle['selected'], bundle['columns']
    missing_selected = [name for name in selected if name not in frame]
    if missing_selected:
        raise ValueError(f'Missing required selected columns: {missing_selected}')
    frame = numeric_columns(frame, [name for name in columns if name in frame], 'prediction input')
    if not frame[selected].isna().any().any():
        probability = predict_complete_selected(artifact, frame)
    else:
        missing_auxiliary = [name for name in columns if name not in frame]
        if missing_auxiliary:
            raise ValueError('A selected predictor has missing values. Supply all 38 candidate columns '
                             'for the fitted imputer. Missing columns: ' + ', '.join(missing_auxiliary))
        probability = predict_saved(artifact, frame)
    if not np.isfinite(probability).all():
        raise RuntimeError('Prediction produced nonfinite probabilities.')
    threshold = float(artifact['threshold'])
    result = pd.DataFrame({'probability_dNCR': probability, 'threshold': threshold,
                           'predicted_class': (probability >= threshold).astype(int)})
    if 'row_id' in frame:
        result.insert(0, 'row_id', frame['row_id'].to_numpy())
    return result


def main():
    package = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True, help='CSV or XLSX; see README.md')
    parser.add_argument('--output', type=Path, required=True, help='New CSV output path')
    parser.add_argument('--model', choices=NAMES, default='LR')
    parser.add_argument('--model-dir', type=Path, default=package / 'model')
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        raise ValueError('Prediction output must differ from the input file.')
    if args.output.exists():
        raise FileExistsError('Choose a new output file; existing files are not overwritten.')
    artifact = joblib.load(args.model_dir / f'{args.model}_pipeline.joblib')
    result = predict_frame(artifact, read_frame(args.input))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    result.to_csv(args.output, index=False)
    print(json.dumps({'model': args.model, 'rows': len(result), 'output': str(args.output)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
